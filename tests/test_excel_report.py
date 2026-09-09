"""Tests the Excel results workbook independent of any LLM call — it never uses one."""
from pathlib import Path

from openpyxl import load_workbook

from tax_agent.calc import run_calculation
from tax_agent.intake.excel_reader import read_intake
from tax_agent.reporting.excel_report import ExcelReportBuilder

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "northern_peak_sample.xlsx"


def test_excel_report_builds_without_error(tmp_path):
    intake = read_intake(FIXTURE_PATH)
    results = run_calculation(intake, tax_year=2025)

    out_path = tmp_path / "Calculation_Results.xlsx"
    ExcelReportBuilder(results).save(out_path)

    assert out_path.exists()
    assert out_path.stat().st_size > 0

    wb = load_workbook(out_path, data_only=True)
    assert wb.sheetnames == [
        "Summary", "SBD", "Corporate Tax", "Tax Pools", "Payroll", "GST-HST", "Flags",
    ]

    summary = wb["Summary"]
    summary_values = [row[1] for row in summary.iter_rows(min_row=4, max_col=2, values_only=True)]
    assert "Northern Peak Consulting Inc." in summary_values
    assert 46568.50 in summary_values  # total tax before credits

    tax_pools = wb["Tax Pools"]
    tax_pools_values = [row[1] for row in tax_pools.iter_rows(min_row=4, max_col=2, values_only=True)]
    assert 22000.0 in tax_pools_values  # CDA balance

    flags = wb["Flags"]
    flag_texts = [row[0] for row in flags.iter_rows(min_row=3, max_col=1, values_only=True)]
    assert any("TOSI risk" in f for f in flag_texts if f)
