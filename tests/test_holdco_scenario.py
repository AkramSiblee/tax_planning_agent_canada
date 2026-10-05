"""
Regression tests for an investment-heavy holding company: Northern Peak Holdings Inc.,
the associated holdco of the Northern Peak Consulting Inc. fixture, for tax year 2026.

It exercises every holdco rule at once: split investment income, Part IV on portfolio
and connected dividends, both RDTOH pools, the dividend refund ordering, GRIP and CDA
additions, and BOTH over-designation penalty taxes. Expected values are hand-computed
below from the 2026 rates workbook (Alberta: general 23%, investment add-on 23.67%).
"""
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from tax_agent.calc import run_calculation
from tax_agent.calc.dividends import compute_dividend_refund
from tax_agent.calc.rates import load_rates
from tax_agent.intake import field_labels as L
from tax_agent.intake.excel_reader import read_intake
from tax_agent.intake.schema import DividendsPaid, Intake, InvestmentIncomeDetail
from tax_agent.reporting.docx_report import DocxReportBuilder
from tax_agent.reporting.excel_report import ExcelReportBuilder


def _holdco_intake(**overrides) -> Intake:
    fields = dict(
        entity_name="Northern Peak Holdings Inc.",
        province="Alberta",
        tax_year=2026,
        is_associated=True,
        net_active_business_income=0,
        aggregate_investment_income=0,  # ignored: investment_detail supersedes it
        net_income_before_tax=375_000,
        group_taxable_capital=1_565_000,
        schedule_23_allocation=150_000,
        non_eligible_rdtoh_opening=12_000,
        eligible_rdtoh_opening=5_000,
        grip_opening=40_000,
        cda_opening=30_000,
        total_salaries=0,
        employee_cpp_total=0,
        employee_cpp2_total=0,
        employee_ei_total=0,
        wcb_assessable_payroll=0,
        wcb_premium_rate_per_100=0,
        taxable_supplies=0,
        total_itcs=0,
        instalments_paid=40_000,
        sred_itc=0,
        provincial_credit=0,
        investment_detail=InvestmentIncomeDetail(
            interest_and_rents=85_000,  # 60,000 bond interest + 25,000 net passive rent
            taxable_capital_gains=90_000,  # 1/2 of a 180,000 gain on a portfolio sale
            eligible_portfolio_dividends=30_000,  # Canadian public shares
            connected_eligible_dividends=50_000,  # from the opco
            connected_non_eligible_dividends=120_000,  # from the opco
            # Opco's refund on those dividends: 18,500 from its eligible RDTOH (its whole
            # balance) and 120,000 x 38 1/3% = 46,000 from its non-eligible RDTOH.
            connected_part_iv_from_payer_eligible_rdtoh=18_500,
            connected_part_iv_from_payer_non_eligible_rdtoh=46_000,
        ),
        dividends_paid=DividendsPaid(eligible=150_000, non_eligible=200_000, capital=150_000),
    )
    fields.update(overrides)
    return Intake(**fields)


class TestHoldcoScenario:
    def setup_method(self):
        self.r = run_calculation(_holdco_intake(), tax_year=2026)

    def test_investment_income_split(self):
        inv = self.r.investment_income
        assert inv.part_i_investment_income == pytest.approx(175_000)  # 85,000 + 90,000
        assert inv.adjusted_aaii_for_grind == pytest.approx(205_000)  # + 30,000 portfolio divs
        assert inv.non_taxable_capital_gains == pytest.approx(90_000)

    def test_sbd_fully_ground(self):
        # (205,000 - 50,000) x 5 = 775,000 -> capped at the full 500,000
        assert self.r.sbd.passive_grind == pytest.approx(500_000)
        assert self.r.sbd.entity_limit == pytest.approx(0)

    def test_part_i_tax_excludes_dividends(self):
        # 175,000 x (23% + 23.67%) = 81,672.50; dividends are s.112-deductible
        assert self.r.corporate_tax.tax_on_investment_income == pytest.approx(81_672.50, abs=0.01)
        assert self.r.corporate_tax.total_tax_before_credits == pytest.approx(81_672.50, abs=0.01)

    def test_part_iv(self):
        inv = self.r.investment_income
        assert inv.part_iv_on_portfolio_dividends == pytest.approx(11_500, abs=0.01)  # 30,000 x 38 1/3%
        assert inv.part_iv_on_connected_dividends == pytest.approx(64_500)
        assert inv.part_iv_tax == pytest.approx(76_000, abs=0.01)
        assert inv.part_iv_to_eligible_rdtoh == pytest.approx(30_000, abs=0.01)  # 11,500 + 18,500
        assert inv.part_iv_to_non_eligible_rdtoh == pytest.approx(46_000)

    def test_rdtoh_pools(self):
        tp = self.r.tax_pools
        # 175,000 x 30.67% = 53,672.50 + 46,000 Part IV = 99,672.50
        assert tp.non_eligible_rdtoh_addition == pytest.approx(99_672.50, abs=0.01)
        assert tp.non_eligible_rdtoh_closing == pytest.approx(111_672.50, abs=0.01)
        assert tp.eligible_rdtoh_closing == pytest.approx(35_000, abs=0.01)

    def test_grip_and_cda_additions(self):
        tp = self.r.tax_pools
        assert tp.grip_addition == pytest.approx(80_000)  # eligible dividends received; no general-rate ABI
        assert tp.grip_available == pytest.approx(120_000)
        assert tp.cda_addition == pytest.approx(90_000)
        assert tp.cda_available == pytest.approx(120_000)

    def test_dividend_refund(self):
        d = self.r.dividends
        # Eligible: 150,000 x 38 1/3% = 57,500, capped at eligible RDTOH 35,000
        assert d.refund_from_eligible_rdtoh == pytest.approx(35_000, abs=0.01)
        # Non-eligible: 200,000 x 38 1/3% = 76,666.67, within non-eligible RDTOH
        assert d.refund_from_non_eligible_rdtoh == pytest.approx(76_666.67, abs=0.01)
        assert d.dividend_refund == pytest.approx(111_666.67, abs=0.01)
        assert d.non_eligible_rdtoh_carryforward == pytest.approx(35_005.83, abs=0.01)
        assert d.eligible_rdtoh_carryforward == pytest.approx(0, abs=0.01)

    def test_penalty_taxes(self):
        d = self.r.dividends
        assert d.excessive_eligible_designation == pytest.approx(30_000)  # 150,000 - 120,000 GRIP
        assert d.part_iii_1_tax == pytest.approx(6_000)  # 20%
        assert d.excessive_capital_dividend == pytest.approx(30_000)  # 150,000 - 120,000 CDA
        assert d.part_iii_tax == pytest.approx(18_000)  # 60%
        assert d.grip_carryforward == pytest.approx(0)
        assert d.cda_carryforward == pytest.approx(0)

    def test_year_end_position(self):
        # 81,672.50 + 76,000 + 6,000 + 18,000 = 181,672.50 before refund
        assert self.r.total_tax_payable_before_refund == pytest.approx(181_672.50, abs=0.01)
        # less 111,666.67 refund = 70,005.83; 40,000 paid -> 30,005.83 owing
        assert self.r.net_corporate_tax_after_credits == pytest.approx(70_005.83, abs=0.01)
        assert self.r.estimated_year_end_position == pytest.approx(-30_005.83, abs=0.01)

    def test_flags(self):
        flags = " | ".join(self.r.flags)
        assert "Excessive eligible dividend designation" in flags
        assert "s.185.1(2)" in flags
        assert "Capital dividend exceeds the CDA" in flags
        assert "s.184(3)" in flags
        assert "RDTOH remains unrecovered" in flags
        assert "prior-year" in flags
        assert "Capital Dividend Account balance is unused" not in flags  # CDA fully used


class TestRefundOrdering:
    def test_non_eligible_dividends_spill_into_eligible_rdtoh(self):
        rates = load_rates(2026)
        from_eligible, from_non_eligible = compute_dividend_refund(
            eligible_paid=0, non_eligible_paid=100_000,
            eligible_rdtoh=50_000, non_eligible_rdtoh=10_000, rates=rates,
        )
        assert from_non_eligible == pytest.approx(10_000)
        assert from_eligible == pytest.approx(28_333.33, abs=0.01)  # 38,333.33 - 10,000

    def test_eligible_dividends_never_touch_non_eligible_rdtoh(self):
        rates = load_rates(2026)
        from_eligible, from_non_eligible = compute_dividend_refund(
            eligible_paid=100_000, non_eligible_paid=0,
            eligible_rdtoh=0, non_eligible_rdtoh=50_000, rates=rates,
        )
        assert from_eligible == pytest.approx(0)
        assert from_non_eligible == pytest.approx(0)


class TestGrindInputs:
    def test_prior_year_aaii_override_drives_the_grind(self):
        r = run_calculation(
            _holdco_intake(prior_year_aaii_override=70_000, net_active_business_income=100_000),
            tax_year=2026,
        )
        # (70,000 - 50,000) x 5 = 100,000 grind, not the current-year 500,000
        assert r.sbd.passive_grind == pytest.approx(100_000)
        assert r.sbd.entity_limit == pytest.approx(150_000)  # Schedule 23 allocation binds
        assert not any("prior-year" in f for f in r.flags)

    def test_general_rate_income_adds_to_grip(self):
        r = run_calculation(
            _holdco_intake(
                net_active_business_income=300_000, prior_year_aaii_override=0,
                schedule_23_allocation=100_000,
            ),
            tax_year=2026,
        )
        # 200,000 above the 100,000 allocation x 0.72 = 144,000 + 80,000 eligible divs received
        assert r.sbd.abi_at_general_rate == pytest.approx(200_000)
        assert r.tax_pools.grip_addition == pytest.approx(224_000)


TEMPLATE_PATH = Path(__file__).parents[1] / "clients" / "_template" / "CCPC_Tax_Planning_Intake_Workbook.xlsx"
FIXTURE_PATH = Path(__file__).parent / "fixtures" / "northern_peak_sample.xlsx"


def _set_by_label(ws, label, value, value_col):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == label:
            ws.cell(row=row, column=value_col, value=value)
            return
    raise AssertionError(f"label {label!r} not in template")


class TestIntakeReaderHoldcoSection:
    def test_fixture_without_dividends_paid(self):
        # The opco fixture fills the investment-detail block but not dividends paid or
        # prior-year AAII — each optional block is independent.
        intake = read_intake(FIXTURE_PATH)
        assert intake.investment_detail is not None
        assert intake.dividends_paid is None
        assert intake.prior_year_aaii_override is None

    def test_template_holdco_section_round_trips(self, tmp_path):
        wb = load_workbook(TEMPLATE_PATH)
        ws = wb["3. Income Statement"]
        _set_by_label(ws, L.INVESTMENT_DETAIL["interest_and_rents"], 85_000, 2)
        _set_by_label(ws, L.INVESTMENT_DETAIL["taxable_capital_gains"], 90_000, 2)
        _set_by_label(ws, L.INVESTMENT_DETAIL["eligible_portfolio_dividends"], 30_000, 2)
        _set_by_label(ws, L.INVESTMENT_DETAIL["connected_part_iv_from_payer_non_eligible_rdtoh"], 46_000, 2)
        _set_by_label(ws, L.DIVIDENDS_PAID["capital"], 150_000, 2)
        _set_by_label(wb["5. Prior-Year Tax Attributes"], L.PRIOR_YEAR_AAII, 70_000, 3)
        path = tmp_path / "intake.xlsx"
        wb.save(path)

        intake = read_intake(path)
        assert intake.investment_detail.interest_and_rents == pytest.approx(85_000)
        assert intake.investment_detail.taxable_capital_gains == pytest.approx(90_000)
        assert intake.investment_detail.eligible_portfolio_dividends == pytest.approx(30_000)
        assert intake.investment_detail.connected_part_iv_from_payer_non_eligible_rdtoh == pytest.approx(46_000)
        assert intake.investment_detail.connected_eligible_dividends == pytest.approx(0)  # blank -> 0
        assert intake.dividends_paid.capital == pytest.approx(150_000)
        assert intake.dividends_paid.eligible == pytest.approx(0)
        assert intake.prior_year_aaii_override == pytest.approx(70_000)

    def test_blank_template_section_falls_back_to_single_aaii(self):
        # A blank optional section must not replace the AAII figure with zeros.
        intake = read_intake(TEMPLATE_PATH)
        assert intake.investment_detail is None
        assert intake.dividends_paid is None
        assert intake.prior_year_aaii_override is None
        results = run_calculation(intake, tax_year=2025)
        assert results.investment_income.part_i_investment_income == pytest.approx(55_000)


class TestHoldcoReports:
    def test_excel_and_docx_build(self, tmp_path):
        results = run_calculation(_holdco_intake(), tax_year=2026)

        xlsx = tmp_path / "results.xlsx"
        ExcelReportBuilder(results).save(xlsx)
        dividends = load_workbook(xlsx)["Dividends"]
        values = [row[1] for row in dividends.iter_rows(min_row=4, max_col=2, values_only=True)]
        assert pytest.approx(18_000) in values  # Part III tax

        docx = tmp_path / "report.docx"
        DocxReportBuilder(results, narrative_text="(test narrative)").save(docx)
        text = "\n".join(
            cell.text for table in Document(docx).tables for row in table.rows for cell in row.cells
        )
        assert "$18,000.00" in text
