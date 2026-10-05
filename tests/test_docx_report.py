"""
Tests the report-building code path independent of any live LLM call — narrative text
is supplied directly, exactly as generate_recommendations() would return it.
"""
from pathlib import Path

from tax_agent.calc import run_calculation
from tax_agent.intake.excel_reader import read_intake
from tax_agent.reporting.docx_report import DocxReportBuilder

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "northern_peak_sample.xlsx"

FAKE_NARRATIVE = """1. Elect a capital dividend now. The $42,000.00 CDA balance can be paid to shareholders tax-free.
2. Review Daniel Okafor's TOSI exposure. His involvement level creates risk under the split-income rules.
3. Monitor the passive-income grind. Aggregate investment income is already reducing the SBD limit."""


def test_docx_report_builds_without_error(tmp_path):
    intake = read_intake(FIXTURE_PATH)
    results = run_calculation(intake, tax_year=2025)

    out_path = tmp_path / "Tax_Planning_Report.docx"
    DocxReportBuilder(results, FAKE_NARRATIVE).save(out_path)

    assert out_path.exists()
    assert out_path.stat().st_size > 0

    # Re-open with python-docx to confirm the file is structurally valid, not just present.
    # Table text lives separately from doc.paragraphs in python-docx's object model —
    # both must be checked to see everything that will actually render in Word.
    from docx import Document
    doc = Document(str(out_path))
    paragraph_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    full_text = paragraph_text + "\n" + table_text

    assert "Northern Peak Consulting Inc." in full_text
    assert "Planning Recommendations" in full_text
    assert "39,568.00" in table_text  # Part I tax before credits appears in the tax-position table
    assert "21,083.33" in table_text  # Part IV tax appears in the executive summary
    assert "42,000.00" in table_text  # CDA balance appears in the tax-pools table
