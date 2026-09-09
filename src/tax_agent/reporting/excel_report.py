"""
Builds a multi-tab Excel workbook of CalculationResults — the same numbers that go into
calculation_results.json and the Word report's tables, laid out for a reviewer who wants
to open the figures directly in Excel rather than parse JSON or read prose.

Contains no narrative and makes no LLM call — this is a pure calc-engine dump, so it is
written on every `run` (even with --skip-narrative). See skills/report-generation/SKILL.md.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from tax_agent.calc.engine import CalculationResults

FONT_NAME = "Arial"
NAVY = "1F3864"
LTBLUE = "D9E2F3"
WHITE = "FFFFFF"

FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
FONT_HEADER = Font(name=FONT_NAME, size=10, bold=True)
FONT_LABEL = Font(name=FONT_NAME, size=10)

FILL_TITLE = PatternFill("solid", fgColor=NAVY)
FILL_HEADER = PatternFill("solid", fgColor=LTBLUE)

THIN = Side(style="thin", color="BFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CURRENCY_FMT = "$#,##0.00;($#,##0.00)"
PERCENT_FMT = "0.00%"


def _title_bar(ws: Worksheet, text: str, span: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    cell = ws.cell(row=1, column=1, value=text)
    cell.font = FONT_TITLE
    cell.fill = FILL_TITLE
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 24


def _write_rows(ws: Worksheet, rows: list[tuple[str, object, str | None]], start_row: int = 3) -> None:
    """Each row is (label, value, number_format). number_format=None leaves it as text/int."""
    headers = ["Item", "Value"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=i, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.border = BORDER_ALL
    row = start_row + 1
    for label, value, number_format in rows:
        ws.cell(row=row, column=1, value=label).font = FONT_LABEL
        vcell = ws.cell(row=row, column=2, value=value)
        vcell.font = FONT_LABEL
        if number_format:
            vcell.number_format = number_format
        ws.cell(row=row, column=1).border = BORDER_ALL
        ws.cell(row=row, column=2).border = BORDER_ALL
        row += 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 24
    ws.freeze_panes = f"A{start_row + 1}"


class ExcelReportBuilder:
    """Writes CalculationResults to a multi-tab .xlsx workbook. No narrative involved."""

    def __init__(self, results: CalculationResults):
        self.results = results

    def _build_summary(self, wb: Workbook) -> None:
        r = self.results
        ws = wb.create_sheet("Summary")
        _title_bar(ws, f"{r.entity_name} — Tax Year {r.tax_year} Summary", span=2)
        _write_rows(ws, [
            ("Entity Name", r.entity_name, None),
            ("Province", r.province, None),
            ("Tax Year", r.tax_year, None),
            ("Net Income Before Tax", r.net_income_before_tax, CURRENCY_FMT),
            ("Estimated Corporate Tax (before credits)", r.corporate_tax.total_tax_before_credits, CURRENCY_FMT),
            ("Net Corporate Tax After Credits", r.net_corporate_tax_after_credits, CURRENCY_FMT),
            ("Estimated Year-End Position (+ = refund)", r.estimated_year_end_position, CURRENCY_FMT),
            ("GST/HST Net Tax", r.gst.net_tax, CURRENCY_FMT),
            ("Instalments Required Next Year", r.instalments_required_next_year, None),
        ])

    def _build_sbd(self, wb: Workbook) -> None:
        r = self.results.sbd
        ws = wb.create_sheet("SBD")
        _title_bar(ws, "Small Business Deduction", span=2)
        _write_rows(ws, [
            ("Passive Income Grind", r.passive_grind, CURRENCY_FMT),
            ("Taxable Capital Grind", r.taxcap_grind, CURRENCY_FMT),
            ("Federal Limit After Grinds", r.federal_limit_after_grinds, CURRENCY_FMT),
            ("Entity SBD Limit", r.entity_limit, CURRENCY_FMT),
            ("Active Business Income at SBD Rate", r.abi_at_sbd_rate, CURRENCY_FMT),
            ("Active Business Income at General Rate", r.abi_at_general_rate, CURRENCY_FMT),
            ("Is Grinding", r.is_grinding, None),
        ])

    def _build_corporate_tax(self, wb: Workbook) -> None:
        r = self.results.corporate_tax
        ws = wb.create_sheet("Corporate Tax")
        _title_bar(ws, "Corporate Tax Breakdown", span=2)
        _write_rows(ws, [
            ("Tax on SBD-Rate Income", r.tax_on_sbd_income, CURRENCY_FMT),
            ("Tax on General-Rate Income", r.tax_on_general_income, CURRENCY_FMT),
            ("Tax on Investment Income", r.tax_on_investment_income, CURRENCY_FMT),
            ("Total Tax Before Credits", r.total_tax_before_credits, CURRENCY_FMT),
            ("Investment Income Rate Used", r.investment_income_rate_used, PERCENT_FMT),
        ])

    def _build_tax_pools(self, wb: Workbook) -> None:
        r = self.results.tax_pools
        ws = wb.create_sheet("Tax Pools")
        _title_bar(ws, "RDTOH, GRIP & CDA", span=2)
        _write_rows(ws, [
            ("Non-Eligible RDTOH Addition", r.non_eligible_rdtoh_addition, CURRENCY_FMT),
            ("Non-Eligible RDTOH Closing Balance", r.non_eligible_rdtoh_closing, CURRENCY_FMT),
            ("Eligible RDTOH Closing Balance", r.eligible_rdtoh_closing, CURRENCY_FMT),
            ("GRIP Available", r.grip_available, CURRENCY_FMT),
            ("CDA Available", r.cda_available, CURRENCY_FMT),
            ("Dividend Needed to Recover Non-Eligible RDTOH", r.dividend_needed_to_recover_non_eligible_rdtoh, CURRENCY_FMT),
            ("Dividend Needed to Recover Eligible RDTOH", r.dividend_needed_to_recover_eligible_rdtoh, CURRENCY_FMT),
        ])

    def _build_payroll(self, wb: Workbook) -> None:
        r = self.results.payroll
        ws = wb.create_sheet("Payroll")
        _title_bar(ws, "Payroll (CPP-CPP2-EI-WCB)", span=2)
        _write_rows(ws, [
            ("Total Salaries", r.total_salaries, CURRENCY_FMT),
            ("Employer CPP", r.employer_cpp, CURRENCY_FMT),
            ("Employer CPP2", r.employer_cpp2, CURRENCY_FMT),
            ("Employer EI", r.employer_ei, CURRENCY_FMT),
            ("Estimated WCB Premium", r.estimated_wcb_premium, CURRENCY_FMT),
        ])

    def _build_gst(self, wb: Workbook) -> None:
        r = self.results.gst
        ws = wb.create_sheet("GST-HST")
        _title_bar(ws, "GST-HST", span=2)
        _write_rows(ws, [
            ("GST/HST Collected", r.gst_collected, CURRENCY_FMT),
            ("Total ITCs", r.total_itcs, CURRENCY_FMT),
            ("Net Tax", r.net_tax, CURRENCY_FMT),
            ("Instalments Likely Required", r.instalments_likely_required, None),
        ])

    def _build_flags(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Flags")
        _title_bar(ws, "Planning Flags", span=1)
        ws.column_dimensions["A"].width = 100
        row = 3
        if not self.results.flags:
            ws.cell(row=row, column=1, value="No flags raised.").font = FONT_LABEL
        for flag in self.results.flags:
            cell = ws.cell(row=row, column=1, value=flag)
            cell.font = FONT_LABEL
            cell.alignment = Alignment(wrap_text=True)
            row += 1

    def build(self) -> Workbook:
        wb = Workbook()
        wb.remove(wb.active)
        self._build_summary(wb)
        self._build_sbd(wb)
        self._build_corporate_tax(wb)
        self._build_tax_pools(wb)
        self._build_payroll(wb)
        self._build_gst(wb)
        self._build_flags(wb)
        return wb

    def save(self, path: str | Path) -> None:
        self.build().save(str(path))
