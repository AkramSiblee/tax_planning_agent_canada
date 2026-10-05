"""
Builds the Word planning report from CalculationResults + the LLM-generated narrative.

See skills/report-generation/SKILL.md for the two-stage pipeline this depends on, and
for why the Limitations section is static text, never LLM-generated.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from tax_agent.calc.engine import CalculationResults
from tax_agent.utils import format_currency

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)

STANDARD_LIMITATIONS_TEXT = [
    "All rates are illustrative approximations — verify current federal and provincial "
    "rates, the SBD threshold, and CPP/EI figures before filing.",
    "The passive-income grind uses this year's investment income as a proxy for the "
    "prior-year AAII test the CRA actually applies — confirm the correct prior-year "
    "figure with your accountant.",
    "Part III.1 and Part III taxes are shown as if no corrective election is filed; "
    "dividend refunds assume the return is filed on time and dividends were paid while "
    "the corporation was a CCPC. Confirm dividend timing against the CDA and GRIP on the "
    "actual payment dates.",
    "This report does not model personal tax outcomes for individual shareholders; a "
    "separate Personal Tax module is planned to complete the salary/dividend and "
    "dividend-recovery analysis.",
    "Nothing in this report should be treated as filed tax advice or a substitute for "
    "your accountant's review and sign-off.",
]


class DocxReportBuilder:
    def __init__(self, results: CalculationResults, narrative_text: str):
        self.results = results
        self.narrative_text = narrative_text
        self.doc = Document()
        self._set_default_font()

    def _set_default_font(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

    def _heading(self, text: str, level: int = 1) -> None:
        h = self.doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = NAVY

    def _note(self, text: str) -> None:
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = GREY

    def _table(self, headers: list[str], rows: list[list[str]]) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
        self.doc.add_paragraph()  # spacing after table

    def _build_title(self) -> None:
        r = self.results
        title = self.doc.add_paragraph()
        run = title.add_run("TAX PLANNING & COMPLIANCE SUMMARY")
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = NAVY

        name = self.doc.add_paragraph()
        run = name.add_run(r.entity_name)
        run.bold = True
        run.font.size = Pt(15)

        sub = self.doc.add_paragraph()
        run = sub.add_run(
            f"Fiscal Year {r.tax_year}  |  {r.province} Canadian-Controlled Private "
            f"Corporation  |  Prepared for internal planning discussion"
        )
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = GREY
        self.doc.add_paragraph()

    def _build_executive_summary(self) -> None:
        r = self.results
        self._heading("Executive Summary")
        self.doc.add_paragraph(
            f"This report summarizes {r.entity_name}'s current-year tax position and "
            f"identifies planning opportunities based on the completed intake for fiscal "
            f"year {r.tax_year}."
        )
        self._table(
            ["Metric", "Amount"],
            [
                ["Net Income Before Tax", format_currency(r.net_income_before_tax)],
                ["Estimated Corporate Tax (before credits)", format_currency(r.corporate_tax.total_tax_before_credits)],
                ["Part IV Tax on Dividends Received", format_currency(r.investment_income.part_iv_tax)],
                ["Dividend Refund", format_currency(r.dividends.dividend_refund)],
                ["Credits Applied (against Part I tax)", format_currency(r.credits.credits_applied)],
                ["Net Corporate Tax After Credits", format_currency(r.net_corporate_tax_after_credits)],
                ["Estimated Year-End Position", format_currency(r.estimated_year_end_position) + (" refund" if r.estimated_year_end_position >= 0 else " owing")],
                ["GST/HST Net Tax", format_currency(r.gst.net_tax)],
            ],
        )

    def _build_tax_position(self) -> None:
        r = self.results
        self._heading("1. Current-Year Corporate Tax Position")
        self._heading("1.1  Small Business Deduction", level=2)
        self.doc.add_paragraph(
            f"This entity's SBD limit is {format_currency(r.sbd.entity_limit)} after "
            f"the passive-income grind ({format_currency(r.sbd.passive_grind)}) and any "
            f"associated-corporation allocation."
        )
        self._table(
            ["Income Component", "Estimated Tax"],
            [
                ["Active business income (SBD rate)", format_currency(r.corporate_tax.tax_on_sbd_income)],
                ["Active business income (general rate)", format_currency(r.corporate_tax.tax_on_general_income)],
                ["Aggregate investment income", format_currency(r.corporate_tax.tax_on_investment_income)],
                ["Total Estimated Corporate Tax", format_currency(r.corporate_tax.total_tax_before_credits)],
            ],
        )

    def _build_tax_pools(self) -> None:
        r = self.results
        self._heading("2. Recoverable Tax Pools — RDTOH, GRIP & CDA")
        self._table(
            ["Pool", "Balance"],
            [
                ["Non-Eligible RDTOH", format_currency(r.tax_pools.non_eligible_rdtoh_closing)],
                ["Eligible RDTOH", format_currency(r.tax_pools.eligible_rdtoh_closing)],
                ["GRIP", format_currency(r.tax_pools.grip_available)],
                ["Capital Dividend Account (CDA)", format_currency(r.tax_pools.cda_available)],
            ],
        )
        d = r.dividends
        if d.eligible_paid or d.non_eligible_paid or d.capital_paid:
            self._heading("2.1  Dividends Paid This Year", level=2)
            self._table(
                ["Item", "Amount"],
                [
                    ["Eligible dividends paid", format_currency(d.eligible_paid)],
                    ["Non-eligible dividends paid", format_currency(d.non_eligible_paid)],
                    ["Capital dividends paid", format_currency(d.capital_paid)],
                    ["Dividend refund", format_currency(d.dividend_refund)],
                    ["Part III.1 tax (excess eligible designation)", format_currency(d.part_iii_1_tax)],
                    ["Part III tax (excess capital dividend)", format_currency(d.part_iii_tax)],
                    ["RDTOH carried forward", format_currency(d.non_eligible_rdtoh_carryforward + d.eligible_rdtoh_carryforward)],
                    ["GRIP carried forward", format_currency(d.grip_carryforward)],
                    ["CDA carried forward", format_currency(d.cda_carryforward)],
                ],
            )

    def _build_payroll_gst(self) -> None:
        r = self.results
        self._heading("3. Payroll & GST/HST Compliance")
        self._table(
            ["Item", "Amount"],
            [
                ["Total Salaries & Wages", format_currency(r.payroll.total_salaries)],
                ["Employer CPP Contribution", format_currency(r.payroll.employer_cpp)],
                ["Employer CPP2 Contribution", format_currency(r.payroll.employer_cpp2)],
                ["Employer EI Contribution", format_currency(r.payroll.employer_ei)],
                ["Estimated WCB Premium", format_currency(r.payroll.estimated_wcb_premium)],
                ["GST/HST Net Tax", format_currency(r.gst.net_tax)],
            ],
        )

    def _build_recommendations(self) -> None:
        self._heading("4. Planning Recommendations")
        for line in self.narrative_text.strip().split("\n"):
            line = line.strip()
            if line:
                self.doc.add_paragraph(line)

    def _build_limitations(self) -> None:
        self._heading("5. Limitations & Next Steps")
        for item in STANDARD_LIMITATIONS_TEXT:
            self.doc.add_paragraph(item, style="List Bullet")
        self._note(
            "Prepared from the completed CCPC Tax Planning Intake Workbook. This is a "
            "planning aid, not a filed return."
        )

    def build(self) -> Document:
        self._build_title()
        self._build_executive_summary()
        self._build_tax_position()
        self._build_tax_pools()
        self._build_payroll_gst()
        self._build_recommendations()
        self._build_limitations()
        return self.doc

    def save(self, path: str | Path) -> None:
        self.build().save(str(path))
