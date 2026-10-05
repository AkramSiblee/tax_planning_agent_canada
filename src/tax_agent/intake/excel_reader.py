"""
Reads a completed CCPC intake workbook into an Intake object.

See skills/intake-collection/SKILL.md for why this uses label-based lookup instead of
fixed cell addresses, and for the required-vs-optional field policy.
"""
from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tax_agent.excel_utils import (
    MAX_SCAN_ROWS,
    find_column_index_by_header,
    find_row_by_partial_match,
    find_value_by_label,
)

from . import field_labels as L
from .schema import DividendsPaid, Intake, InvestmentIncomeDetail, MissingIntakeDataError

logger = logging.getLogger(__name__)


def _read_entity_setup(ws: Worksheet) -> dict:
    is_associated_raw = find_value_by_label(ws, L.ENTITY_SETUP["is_associated"], value_col=3)
    return {
        "entity_name": find_value_by_label(ws, L.ENTITY_SETUP["entity_name"], value_col=3),
        "province": find_value_by_label(ws, L.ENTITY_SETUP["province"], value_col=3),
        "tax_year": int(find_value_by_label(ws, L.ENTITY_SETUP["tax_year"], value_col=3)),
        "is_associated": str(is_associated_raw).strip().lower() == "yes",
    }


def _read_income_statement(ws: Worksheet) -> dict:
    labels = L.INCOME_STATEMENT
    return {
        "net_active_business_income": find_value_by_label(ws, labels["net_active_business_income"], value_col=2) or 0.0,
        "aggregate_investment_income": find_value_by_label(ws, labels["aggregate_investment_income"], value_col=2) or 0.0,
        "net_income_before_tax": find_value_by_label(ws, labels["net_income_before_tax"], value_col=2) or 0.0,
    }


def _read_optional_amounts(ws: Worksheet, labels: dict[str, str], value_col: int) -> dict[str, float] | None:
    """
    Reads every label in `labels`. Returns None if the section is absent OR left entirely
    blank/zero, so an unused optional section never overrides the single-figure inputs.
    """
    raw = {
        key: find_value_by_label(ws, label, value_col=value_col, required=False, default=None)
        for key, label in labels.items()
    }
    if not any(raw.values()):
        return None
    return {key: float(value or 0.0) for key, value in raw.items()}


def _read_investment_detail(ws: Worksheet) -> InvestmentIncomeDetail | None:
    amounts = _read_optional_amounts(ws, L.INVESTMENT_DETAIL, value_col=2)
    return None if amounts is None else InvestmentIncomeDetail(**amounts)


def _read_dividends_paid(ws: Worksheet) -> DividendsPaid | None:
    amounts = _read_optional_amounts(ws, L.DIVIDENDS_PAID, value_col=2)
    return None if amounts is None else DividendsPaid(**amounts)


def _read_prior_year_aaii(ws: Worksheet) -> float | None:
    """None when the label is absent OR left blank — blank means 'unknown', not zero."""
    value = find_value_by_label(ws, L.PRIOR_YEAR_AAII, value_col=3, required=False, default=None)
    return None if value in (None, "") else float(value)


def _read_prior_year_attributes(ws: Worksheet) -> dict:
    labels = L.PRIOR_YEAR_ATTRIBUTES
    return {
        "eligible_rdtoh_opening": find_value_by_label(ws, labels["eligible_rdtoh_opening"], value_col=3, required=False) or 0.0,
        "non_eligible_rdtoh_opening": find_value_by_label(ws, labels["non_eligible_rdtoh_opening"], value_col=3, required=False) or 0.0,
        "grip_opening": find_value_by_label(ws, labels["grip_opening"], value_col=3, required=False) or 0.0,
        "cda_opening": find_value_by_label(ws, labels["cda_opening"], value_col=3, required=False) or 0.0,
    }


def _read_associated_corporations(ws: Worksheet, entity_name: str, is_associated: bool) -> dict:
    if not is_associated:
        return {"group_taxable_capital": 0.0, "schedule_23_allocation": None}

    total_row = None
    for row in range(1, MAX_SCAN_ROWS):
        if ws.cell(row=row, column=1).value == "TOTAL":
            total_row = row
            break
    if total_row is None:
        logger.warning(
            "Entity Setup says associated=Yes but no TOTAL row found on "
            "'7. Associated Corporations' — treating group taxable capital as 0."
        )
        return {"group_taxable_capital": 0.0, "schedule_23_allocation": None}

    group_taxable_capital = ws.cell(row=total_row, column=5).value or 0.0

    entity_row = find_row_by_partial_match(ws, column=1, substring=entity_name)
    if entity_row is None:
        logger.warning(
            "Could not find a row on '7. Associated Corporations' matching entity name "
            "%r — this entity's Schedule 23 allocation will be treated as unknown, and "
            "the full grind-adjusted federal SBD limit will apply instead. Confirm the "
            "corporation's own row is present and named consistently with Entity Setup.",
            entity_name,
        )
        return {"group_taxable_capital": group_taxable_capital, "schedule_23_allocation": None}

    schedule_23_allocation = ws.cell(row=entity_row, column=6).value or 0.0
    return {
        "group_taxable_capital": group_taxable_capital,
        "schedule_23_allocation": schedule_23_allocation,
    }


def _read_payroll_register(ws: Worksheet) -> dict:
    total_row = None
    for row in range(1, MAX_SCAN_ROWS):
        if ws.cell(row=row, column=1).value == "TOTAL":
            total_row = row
            break
    if total_row is None:
        raise MissingIntakeDataError("Could not find TOTAL row on '9. Payroll Register'.")

    return {
        "total_salaries": ws.cell(row=total_row, column=3).value or 0.0,
        "employee_cpp_total": ws.cell(row=total_row, column=4).value or 0.0,
        "employee_cpp2_total": ws.cell(row=total_row, column=5).value or 0.0,
        "employee_ei_total": ws.cell(row=total_row, column=6).value or 0.0,
        "wcb_assessable_payroll": find_value_by_label(ws, "WCB Assessable Payroll", value_col=3, required=False) or 0.0,
        "wcb_premium_rate_per_100": find_value_by_label(ws, "WCB Premium Rate (per $100 of assessable payroll)", value_col=3, required=False) or 0.0,
    }


def _read_gst_hst(ws: Worksheet) -> dict:
    labels = L.GST_HST
    return {
        "taxable_supplies": find_value_by_label(ws, labels["taxable_supplies"], value_col=3, required=False) or 0.0,
        "total_itcs": find_value_by_label(ws, labels["total_itcs"], value_col=3, required=False) or 0.0,
    }


def _read_instalments(ws: Worksheet) -> dict:
    return {
        "instalments_paid": find_value_by_label(ws, L.INSTALMENTS["instalments_paid"], value_col=2, required=False) or 0.0,
    }


def _read_credits(ws: Worksheet) -> dict:
    labels = L.CREDITS
    return {
        "sred_itc": find_value_by_label(ws, labels["sred_itc"], value_col=3, required=False) or 0.0,
        "provincial_credit": find_value_by_label(ws, labels["provincial_credit"], value_col=3, required=False) or 0.0,
        "sred_itc_refundable": find_value_by_label(ws, labels["sred_itc_refundable"], value_col=3, required=False) or 0.0,
    }


def _read_tosi_risk_shareholders(ws: Worksheet) -> list[str]:
    header_row = None
    for row in range(1, MAX_SCAN_ROWS):
        if ws.cell(row=row, column=1).value == L.SHAREHOLDERS_TOSI["name_column_header"]:
            header_row = row
            break
    if header_row is None:
        return []

    tosi_col = find_column_index_by_header(ws, header_row, L.SHAREHOLDERS_TOSI["tosi_status_column_header"])
    if tosi_col is None:
        return []

    flagged = []
    for row in range(header_row + 1, MAX_SCAN_ROWS):
        name = ws.cell(row=row, column=1).value
        if name == "TOTAL" or name is None:
            break
        status = ws.cell(row=row, column=tosi_col).value
        if status == L.SHAREHOLDERS_TOSI["not_excluded_value"]:
            flagged.append(name)
    return flagged


def read_intake(workbook_path: str | Path) -> Intake:
    from tax_agent.excel_utils import LabelNotFoundError

    wb = load_workbook(workbook_path, data_only=True)

    try:
        entity = _read_entity_setup(wb["1. Entity Setup"])
        income = _read_income_statement(wb["3. Income Statement"])
    except LabelNotFoundError as exc:
        raise MissingIntakeDataError(str(exc)) from exc

    attrs = _read_prior_year_attributes(wb["5. Prior-Year Tax Attributes"])
    prior_year_aaii = _read_prior_year_aaii(wb["5. Prior-Year Tax Attributes"])
    investment_detail = _read_investment_detail(wb["3. Income Statement"])
    dividends_paid = _read_dividends_paid(wb["3. Income Statement"])
    assoc = _read_associated_corporations(
        wb["7. Associated Corporations"], entity["entity_name"], entity["is_associated"]
    )
    tosi_risk = _read_tosi_risk_shareholders(wb["8. Shareholders & TOSI"])
    payroll = _read_payroll_register(wb["9. Payroll Register"])
    gst = _read_gst_hst(wb["10. GST-HST"])
    instalments = _read_instalments(wb["11. Instalments & Compliance"])
    credits_ = _read_credits(wb["12. Credits & Incentives"])

    return Intake(
        entity_name=entity["entity_name"],
        province=entity["province"],
        tax_year=entity["tax_year"],
        is_associated=entity["is_associated"],
        net_active_business_income=income["net_active_business_income"],
        aggregate_investment_income=income["aggregate_investment_income"],
        net_income_before_tax=income["net_income_before_tax"],
        group_taxable_capital=assoc["group_taxable_capital"],
        schedule_23_allocation=assoc["schedule_23_allocation"],
        non_eligible_rdtoh_opening=attrs["non_eligible_rdtoh_opening"],
        eligible_rdtoh_opening=attrs["eligible_rdtoh_opening"],
        grip_opening=attrs["grip_opening"],
        cda_opening=attrs["cda_opening"],
        total_salaries=payroll["total_salaries"],
        employee_cpp_total=payroll["employee_cpp_total"],
        employee_cpp2_total=payroll["employee_cpp2_total"],
        employee_ei_total=payroll["employee_ei_total"],
        wcb_assessable_payroll=payroll["wcb_assessable_payroll"],
        wcb_premium_rate_per_100=payroll["wcb_premium_rate_per_100"],
        taxable_supplies=gst["taxable_supplies"],
        total_itcs=gst["total_itcs"],
        instalments_paid=instalments["instalments_paid"],
        sred_itc=credits_["sred_itc"],
        provincial_credit=credits_["provincial_credit"],
        sred_itc_refundable=credits_["sred_itc_refundable"],
        tosi_risk_shareholders=tosi_risk,
        prior_year_aaii_override=prior_year_aaii,
        investment_detail=investment_detail,
        dividends_paid=dividends_paid,
    )
