"""
Loads config/Tax_Rates_<year>.xlsx — the single source of truth for every rate,
threshold, and payroll figure the calc engine uses.

See skills/rate-table-maintenance/SKILL.md for how this file gets updated (by whoever
maintains the agent, or an LLM/coding assistant asked to refresh rates — never by the
end client, and never by hardcoding a value into this module or any calc/ function).
"""
from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from tax_agent.excel_utils import find_value_by_label, iter_table_rows

CONFIG_DIR = Path(__file__).resolve().parents[3] / "config"


class UnknownProvinceError(ValueError):
    pass


class RateFileNotFoundError(FileNotFoundError):
    pass


@dataclass(frozen=True)
class ProvinceRates:
    name: str
    general_rate_combined: float
    small_business_rate_combined: float
    sbd_threshold: float
    gst_hst_type: str
    gst_hst_rate: float
    notes: str


@dataclass(frozen=True)
class RateTable:
    tax_year: int
    federal: dict
    payroll: dict
    provinces: dict[str, ProvinceRates]

    def province(self, name: str) -> ProvinceRates:
        try:
            return self.provinces[name]
        except KeyError as exc:
            valid = ", ".join(sorted(self.provinces))
            raise UnknownProvinceError(
                f"{name!r} is not a recognized province/territory. Valid options: {valid}"
            ) from exc


# Every federal-figure name expected on the "Federal & Payroll" tab. Split into the two
# groups the calc engine consumes (rates.federal[...] vs rates.payroll[...]) — this
# mapping is the only place that split is encoded.
_FEDERAL_KEYS = [
    "general_rate_net", "small_business_rate", "sbd_base_limit",
    "investment_income_addon", "rdtoh_addition_rate", "dividend_refund_rate",
    "instalment_threshold", "taxable_capital_grind_start", "taxable_capital_grind_end",
    "passive_income_grind_start", "passive_income_grind_end",
    "part_iv_tax_rate", "capital_gains_inclusion_rate", "grip_general_rate_factor",
    "part_iii_1_rate", "part_iii_rate",
]
_PAYROLL_KEYS = [
    "cpp_employee_max", "cpp2_employee_max", "cpp2_lower_threshold",
    "cpp2_upper_threshold", "ei_employee_max", "ei_employer_multiplier",
]


@lru_cache(maxsize=None)
def load_rates(tax_year: int) -> RateTable:
    path = CONFIG_DIR / f"Tax_Rates_{tax_year}.xlsx"
    if not path.exists():
        available = sorted(p.stem for p in CONFIG_DIR.glob("Tax_Rates_*.xlsx"))
        raise RateFileNotFoundError(
            f"No rates workbook for tax year {tax_year} at {path}. "
            f"Available: {', '.join(available) or '(none found)'}. "
            f"See skills/rate-table-maintenance/SKILL.md to add one."
        )
    wb = load_workbook(path, data_only=True)

    fp_ws = wb["Federal & Payroll"]
    federal = {
        key: find_value_by_label(fp_ws, key, value_col=2) for key in _FEDERAL_KEYS
    }
    payroll = {
        key: find_value_by_label(fp_ws, key, value_col=2) for key in _PAYROLL_KEYS
    }

    prov_ws = wb["Provinces"]
    header_row = 3  # "Province" header — see build_rates_workbook.py layout
    provinces: dict[str, ProvinceRates] = {}
    for row in iter_table_rows(prov_ws, header_row, stop_values=()):
        name = prov_ws.cell(row=row, column=1).value
        if name is None:
            break
        provinces[name] = ProvinceRates(
            name=name,
            general_rate_combined=prov_ws.cell(row=row, column=2).value,
            small_business_rate_combined=prov_ws.cell(row=row, column=3).value,
            sbd_threshold=prov_ws.cell(row=row, column=4).value,
            gst_hst_type=prov_ws.cell(row=row, column=5).value,
            gst_hst_rate=prov_ws.cell(row=row, column=6).value,
            notes=prov_ws.cell(row=row, column=7).value or "",
        )

    return RateTable(
        tax_year=tax_year,
        federal=federal,
        payroll=payroll,
        provinces=provinces,
    )
