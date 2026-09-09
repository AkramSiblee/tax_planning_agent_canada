"""
Shared label-based Excel lookup helpers.

Used by both intake/excel_reader.py (reading a client's intake workbook) and
calc/rates.py (reading the tax rates workbook) — see skills/intake-collection/SKILL.md
and skills/rate-table-maintenance/SKILL.md for why label-based lookup, not fixed cell
addresses, is the rule for every Excel file this agent reads.
"""
from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

MAX_SCAN_ROWS = 300
MAX_SCAN_COLS = 20


class LabelNotFoundError(ValueError):
    pass


def find_value_by_label(
    ws: Worksheet, label: str, value_col: int, *, required: bool = True, default=0.0
):
    """
    Scan column A of `ws` for an exact match of `label`, return the value in
    `value_col` of that row. Raises LabelNotFoundError if required and not found;
    otherwise returns `default`.
    """
    for row in range(1, MAX_SCAN_ROWS):
        cell_value = ws.cell(row=row, column=1).value
        if cell_value == label:
            return ws.cell(row=row, column=value_col).value
    if required:
        raise LabelNotFoundError(f"Could not find label {label!r} on sheet {ws.title!r}.")
    return default


def find_row_by_partial_match(ws: Worksheet, column: int, substring: str) -> int | None:
    """Return the row number whose `column` value contains `substring` (case-insensitive), or None."""
    needle = substring.strip().lower()
    for row in range(1, MAX_SCAN_ROWS):
        cell_value = ws.cell(row=row, column=column).value
        if isinstance(cell_value, str) and needle in cell_value.lower():
            return row
    return None


def find_column_index_by_header(ws: Worksheet, header_row: int, header_text: str) -> int | None:
    for col in range(1, MAX_SCAN_COLS):
        if ws.cell(row=header_row, column=col).value == header_text:
            return col
    return None


def find_header_row(ws: Worksheet, header_text: str, *, search_col: int = 1) -> int | None:
    """Return the row number where `search_col`'s value equals `header_text` exactly."""
    for row in range(1, MAX_SCAN_ROWS):
        if ws.cell(row=row, column=search_col).value == header_text:
            return row
    return None


def iter_table_rows(ws: Worksheet, header_row: int, *, stop_values: tuple[str, ...] = ("TOTAL",)):
    """
    Yield row numbers below `header_row` until a blank first cell or a value in
    `stop_values` is hit. Used to read every data row of a table without knowing in
    advance how many rows it has.
    """
    row = header_row + 1
    while row < MAX_SCAN_ROWS:
        first_cell = ws.cell(row=row, column=1).value
        if first_cell is None or first_cell in stop_values:
            return
        yield row
        row += 1
